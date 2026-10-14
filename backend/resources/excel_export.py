import io
from flask import send_file
from flask_restful import Resource
from ..models import ImportedNumber, EyeconResult, FacebookResult
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

class ExcelExport(Resource):
    def get(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "EyeConExcel Results"

        # Styles
        header_font = Font(bold=True, color="FFFFFF", size=11, name="Calibri")
        header_fill = PatternFill("solid", fgColor="1a1a2e")
        alt_fill = PatternFill("solid", fgColor="16213e")
        border = Border(
            left=Side(style='thin', color='334155'),
            right=Side(style='thin', color='334155'),
            top=Side(style='thin', color='334155'),
            bottom=Side(style='thin', color='334155')
        )

        headers = ["#", "Phone Number", "Eyecon Names", "Facebook ID", "Facebook Link", "Status"]
        col_widths = [5, 20, 45, 25, 45, 12]

        for i, (h, w) in enumerate(zip(headers, col_widths), 1):
            cell = ws.cell(row=1, column=i, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 22

        nums = ImportedNumber.query.order_by(ImportedNumber.id).all()
        for idx, n in enumerate(nums, 2):
            er = EyeconResult.query.filter_by(number=n.number).first()
            fr = FacebookResult.query.filter_by(number=n.number).first()
            fill = alt_fill if idx % 2 == 0 else PatternFill("solid", fgColor="0f3460")
            status = "Searched" if n.searched else "Pending"
            row_data = [
                idx - 1,
                n.number,
                er.name if er else "",
                fr.fb_id if (fr and fr.fb_id) else "",
                fr.fb_link if (fr and fr.fb_link) else "",
                status
            ]
            for col, val in enumerate(row_data, 1):
                cell = ws.cell(row=idx, column=col, value=val)
                cell.fill = fill
                cell.border = border
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                cell.font = Font(color="E2E8F0", name="Calibri")
            ws.row_dimensions[idx].height = 18

        ws.freeze_panes = 'A2'

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name='eyeconexcel_results.xlsx'
        )
